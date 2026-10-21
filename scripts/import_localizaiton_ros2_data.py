#!/usr/bin/env python3
"""Convert the legacy localizaiton_ros2 XLSX files into authoritative CSVs.

This is an upstream, one-time representation conversion.  It never scales or
offsets time, reorders/drops/adds waypoints, or changes XYZ values.  Once a CSV
is written, a different regeneration is refused instead of overwriting it.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import math
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_EVEN
from pathlib import Path
from typing import Iterable, Mapping

import numpy as np
import yaml
from openpyxl import load_workbook
from scipy.spatial.transform import Rotation


SOURCE_REPOSITORY = "https://github.com/hxhou0207/localizaiton_ros2.git"
SOURCE_BRANCH = "ros2"
SOURCE_COMMIT = "cde2f051f610028182721803195c382758e06844"
SOURCE_PATH = "src/moveit_python_test/data"
SOURCE_PROVENANCE = {
    "repository": SOURCE_REPOSITORY,
    "branch": SOURCE_BRANCH,
    "commit": SOURCE_COMMIT,
    "path": SOURCE_PATH,
}
EXPECTED_SOURCE_SHA256 = {
    "3comb_optical_table.xlsx": (
        "5451ba1775936d52137b101142e6877e6c8a999b2cc8ed2111ad815f731b6a63"
    ),
    "8omega3_desk.xlsx": (
        "3ddf52c719ddd35240c777e73656b9a7c64b37b8288574c52b359711d719549a"
    ),
    "comb_desk.xlsx": (
        "ef470289f8312404dbd31b18ee28ed97fb8d509476d0835195aa9994d3cb08b6"
    ),
    "comb_optical_table.xlsx": (
        "6644257a08107b688c0b83399d38db3ed05ad3226106eb14f7d423e322fc03c9"
    ),
    "inner_sphere.xlsx": (
        "c00e1431a823ae55172ee3542b305749c4bd2a6c09e14da8858b0eae144c0202"
    ),
    "omega1.xlsx": (
        "13e0f67ce31f284ace96badbd938009c6386b26989067593756f03ba20017b0d"
    ),
    "omega3.xlsx": (
        "524a51eb5d8248a162104dff7d4a6c198455128d5fd875ff12a9895bf049bae4"
    ),
    "omega3_desk.xlsx": (
        "9f35e182ab8c63cac475aa70600383c89fe69cb769f1bc5cf0d8fe7b97501b54"
    ),
    "outer_sphere.xlsx": (
        "40baa547a69136545ac4c04ff29b72d18c4934ada0abea5fc13503ada291909d"
    ),
    "snake_quat_fixedroll.xlsx": (
        "5008885cba3a3cb1dee037d063aafb242425d9f2d1cf2cea2ac8280fe9408f99"
    ),
    "snake_quat_fixedroll_2.xlsx": (
        "0f222dfdefafc47a327e8f81a87916cab1e6aebf65b985dda37a0a6ec70017c9"
    ),
    "sphere_optical_table.xlsx": (
        "18d38cfeac5e85e3d9e3867d222380a453da300194032bdc5a206cb050c86b3a"
    ),
    "trajectory_circle_desk.xlsx": (
        "373b7dd8a0a6153087f7b604f6603752e8ec7ceb87d13d3e605ba98825bef39c"
    ),
    "trajectory_circle_optical_table.xlsx": (
        "c2a4d8849278e495eb3afa25dd26265817d817ee3bbeadc59e8fdefae2d07a8a"
    ),
    "trajectory_snake3d_desk.xlsx": (
        "1233f30a81976f3bc0ecac8ef7940ffecbb372cadfffb017eaadc41899b8646f"
    ),
    "trajectory_snake3d_optical_table.xlsx": (
        "450b0e3fe50ee3463819c5e521acc8c392591313f00056d7980e6af87a7dc851"
    ),
    "trajectory_snake3d_optical_table_part.xlsx": (
        "de2c296ca6264948550f22b615448a810301194d624f15a086ddd37da8a2e01b"
    ),
    "trajectory_sphere_desk.xlsx": (
        "80ad4fb33def0f8fdd9e8dbd64969098fd6c2f45debe9ae06a49484a8fa10b43"
    ),
    "trajectory_square_desk.xlsx": (
        "eb14c4c0e14d6dce128669ccda96a7f431936cdc580d7ee4e11586d935aa514f"
    ),
    "trajectory_square_optical_table.xlsx": (
        "174420638f3417969848e3695898cf81c55e7638c2f3ecd96e1fc883b06523a2"
    ),
}
PLANE_POSE_REGISTRY_PATH = (
    Path(__file__).resolve().parents[1]
    / "environments"
    / "legacy_plane1_pose_registry.yaml"
)
CSV_HEADER = (
    "time_s",
    "frame_id",
    "tip_x_m",
    "tip_y_m",
    "tip_z_m",
    "tip_qx",
    "tip_qy",
    "tip_qz",
    "tip_qw",
    "phase",
)


@dataclass(frozen=True)
class TeachPendantPlanePose:
    pose_id: str
    source_row: int
    raw_base_to_plane1_pose_xyz_rotvec: tuple[float, ...]
    translation_m: tuple[float, ...]
    quaternion_xyzw: tuple[float, ...]
    plane_z_axis_base_link: tuple[float, ...]
    normal_direction: str
    exact_duplicate_of: str | None


def _registry_vector(
    value: object,
    size: int,
    label: str,
) -> tuple[float, ...]:
    if not isinstance(value, list) or len(value) != size:
        raise ValueError(f"{label} must contain exactly {size} values")
    if any(isinstance(item, bool) for item in value):
        raise ValueError(f"{label} must contain only finite numbers")
    try:
        parsed = tuple(float(item) for item in value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label} must contain only finite numbers") from error
    if any(not math.isfinite(item) for item in parsed):
        raise ValueError(f"{label} must contain only finite numbers")
    return parsed


def load_plane_pose_registry(
    path: Path = PLANE_POSE_REGISTRY_PATH,
) -> dict[str, TeachPendantPlanePose]:
    try:
        document = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as error:
        raise ValueError(f"cannot load plane pose registry {path}: {error}") from error
    if not isinstance(document, dict) or document.get("schema_version") != 1:
        raise ValueError("plane pose registry must be a schema_version 1 mapping")

    source = document.get("source")
    if not isinstance(source, dict) or source.get("supplied_row_count") != 16:
        raise ValueError("plane pose registry must declare all 16 supplied source rows")
    conversion = document.get("base_link_conversion")
    if not isinstance(conversion, dict):
        raise ValueError("plane pose registry is missing base_link_conversion")
    if conversion.get("conversion_id") != "controlled_urdf_base_link_to_base_rz_pi":
        raise ValueError("plane pose registry uses an unsupported base_link conversion")
    conversion_translation = _registry_vector(
        conversion.get("translation_m"), 3, "base_link conversion translation"
    )
    conversion_rpy = _registry_vector(
        conversion.get("rotation_rpy_rad"), 3, "base_link conversion RPY"
    )
    if not np.allclose(conversion_translation, (0.0, 0.0, 0.0), atol=1.0e-15):
        raise ValueError("base_link conversion translation must be zero")
    if not np.allclose(
        conversion_rpy,
        (0.0, 0.0, math.pi),
        rtol=0.0,
        atol=1.0e-15,
    ):
        raise ValueError("base_link conversion must be exactly Rz(pi)")
    base_link_to_base = Rotation.from_euler("xyz", conversion_rpy)

    raw_poses = document.get("poses")
    if not isinstance(raw_poses, dict) or len(raw_poses) != 16:
        raise ValueError("plane pose registry must contain exactly 16 pose mappings")
    poses: dict[str, TeachPendantPlanePose] = {}
    for pose_id, value in raw_poses.items():
        if not isinstance(pose_id, str) or not isinstance(value, dict):
            raise ValueError("plane pose registry entries must be named mappings")
        source_row = value.get("source_row")
        if isinstance(source_row, bool) or not isinstance(source_row, int):
            raise ValueError(f"{pose_id}.source_row must be an integer")
        if pose_id != f"user_measured_candidate_p{source_row}":
            raise ValueError(f"{pose_id} does not match source row {source_row}")
        raw_pose = _registry_vector(
            value.get("raw_base_to_plane1_pose_xyz_rotvec"),
            6,
            f"{pose_id}.raw_base_to_plane1_pose_xyz_rotvec",
        )
        derived = value.get("derived_base_link_to_plane1")
        if not isinstance(derived, dict):
            raise ValueError(f"{pose_id} is missing its derived base_link transform")
        translation = _registry_vector(
            derived.get("translation_m"), 3, f"{pose_id}.translation_m"
        )
        quaternion = _registry_vector(
            derived.get("quaternion_xyzw"), 4, f"{pose_id}.quaternion_xyzw"
        )
        plane_z_axis = _registry_vector(
            derived.get("plane_z_axis_base_link"),
            3,
            f"{pose_id}.plane_z_axis_base_link",
        )
        if quaternion[3] < 0.0:
            raise ValueError(f"{pose_id} quaternion must use the canonical qw >= 0 sign")

        expected_translation = base_link_to_base.apply(raw_pose[:3])
        expected_rotation = base_link_to_base * Rotation.from_rotvec(raw_pose[3:])
        expected_quaternion = expected_rotation.as_quat()
        if expected_quaternion[3] < 0.0:
            expected_quaternion *= -1.0
        expected_plane_z_axis = expected_rotation.apply((0.0, 0.0, 1.0))
        if not np.allclose(translation, expected_translation, rtol=0.0, atol=1.0e-12):
            raise ValueError(f"{pose_id} derived translation does not match Rz(pi)")
        if not np.allclose(quaternion, expected_quaternion, rtol=0.0, atol=1.0e-12):
            raise ValueError(f"{pose_id} derived quaternion does not match Rz(pi)")
        if not np.allclose(plane_z_axis, expected_plane_z_axis, rtol=0.0, atol=1.0e-12):
            raise ValueError(f"{pose_id} plane Z-axis evidence does not match quaternion")
        expected_direction = (
            "base_link_negative_z"
            if expected_plane_z_axis[2] < 0.0
            else "base_link_positive_z"
        )
        if value.get("normal_direction") != expected_direction:
            raise ValueError(f"{pose_id} normal_direction does not match its transform")
        exact_duplicate_of = value.get("exact_duplicate_of")
        if exact_duplicate_of is not None and not isinstance(exact_duplicate_of, str):
            raise ValueError(f"{pose_id}.exact_duplicate_of must be a pose ID")
        poses[pose_id] = TeachPendantPlanePose(
            pose_id=pose_id,
            source_row=source_row,
            raw_base_to_plane1_pose_xyz_rotvec=raw_pose,
            translation_m=translation,
            quaternion_xyzw=quaternion,
            plane_z_axis_base_link=plane_z_axis,
            normal_direction=expected_direction,
            exact_duplicate_of=exact_duplicate_of,
        )

    if {pose.source_row for pose in poses.values()} != set(range(1, 17)):
        raise ValueError("plane pose registry source rows must be exactly 1 through 16")
    for pose in poses.values():
        if pose.exact_duplicate_of is None:
            continue
        original = poses.get(pose.exact_duplicate_of)
        if original is None:
            raise ValueError(f"{pose.pose_id} references an unknown duplicate source")
        if (
            pose.raw_base_to_plane1_pose_xyz_rotvec
            != original.raw_base_to_plane1_pose_xyz_rotvec
            or pose.translation_m != original.translation_m
            or pose.quaternion_xyzw != original.quaternion_xyzw
        ):
            raise ValueError(f"{pose.pose_id} is not an exact duplicate of its target")
    duplicate_rows = {
        pose.pose_id: pose.exact_duplicate_of
        for pose in poses.values()
        if pose.exact_duplicate_of is not None
    }
    if duplicate_rows != {
        "user_measured_candidate_p14": "user_measured_candidate_p8"
    }:
        raise ValueError("the supplied duplicate evidence must be exactly p14 -> p8")
    return poses


PLANE_POSES = load_plane_pose_registry()


@dataclass(frozen=True)
class LegacyProfile:
    profile_id: str
    environment_id: str
    rod_length_m: float
    plane_pose_id: str
    legacy_executor: str
    legacy_tcp_evidence: str
    selection_basis: str


TEST_PROFILES = {
    "legacy_plane1_rod217": LegacyProfile(
        profile_id="p15_rod217_desk",
        environment_id="legacy_plane1_rod217",
        rod_length_m=0.217,
        plane_pose_id="user_measured_candidate_p15",
        legacy_executor="play_excel_trajectory",
        legacy_tcp_evidence=(
            "active 0.217 m default at source commit; no per-XLSX CLI binding survives"
        ),
        selection_basis=(
            "fixed desk-family profile: the near-field p15 candidate and horizontal "
            "tool axis make the longer rod move tool0 toward the robot base"
        ),
    ),
    "legacy_plane1_test_rod120": LegacyProfile(
        profile_id="p15_rod120_optical",
        environment_id="legacy_plane1_test_rod120",
        rod_length_m=0.120,
        plane_pose_id="user_measured_candidate_p15",
        legacy_executor="play_excel_trajectory",
        legacy_tcp_evidence=(
            "Z-vector schema is accepted by the legacy executor; 0.120 m is an "
            "explicit standard test-rod assumption rather than a recovered file binding"
        ),
        selection_basis=(
            "fixed ordinary optical/sphere profile: shorter rod lowers the wrist radius "
            "for trajectories whose tool Z axis points approximately along -plane Z"
        ),
    ),
    "legacy_plane1_test_p10_rod120": LegacyProfile(
        profile_id="p10_rod120_large_optical",
        environment_id="legacy_plane1_test_p10_rod120",
        rod_length_m=0.120,
        plane_pose_id="user_measured_candidate_p10",
        legacy_executor="play_excel_trajectory",
        legacy_tcp_evidence=(
            "Z-vector schema is accepted by the legacy executor; 0.120 m is the same "
            "standard test rod used by the ordinary optical profile"
        ),
        selection_basis=(
            "fixed large-envelope optical profile: the measured p10 candidate is closer "
            "to the robot base and is selected for the complete snake family"
        ),
    ),
    "legacy_plane1_p16_rod194": LegacyProfile(
        profile_id="p16_rod194_quaternion_flipped_normal",
        environment_id="legacy_plane1_p16_rod194",
        rod_length_m=0.194,
        plane_pose_id="user_measured_candidate_p16",
        legacy_executor="go_to_point",
        legacy_tcp_evidence=(
            "active 0.194 m default in the quaternion/magnet-map revision; a same-day "
            "snake_quat_fixedroll map survives, but no per-file plane pose binding does"
        ),
        selection_basis=(
            "fixed quaternion-snake profile: preserve the strongest legacy rod evidence; "
            "p16 is the candidate selected in the user's current conversion script and "
            "its plane Z normal is flipped relative to p1-p15"
        ),
    ),
}
if unknown_pose_ids := sorted(
    profile.plane_pose_id
    for profile in TEST_PROFILES.values()
    if profile.plane_pose_id not in PLANE_POSES
):
    raise RuntimeError(
        "legacy profile references an unknown teach-pendant pose: "
        + ", ".join(unknown_pose_ids)
    )


_TRAJECTORY_PROFILE_ASSIGNMENTS = (
    ("omega3", "legacy_plane1_test_rod120"),
    ("omega3_desk", "legacy_plane1_rod217"),
    ("8omega3_desk", "legacy_plane1_rod217"),
    ("omega1", "legacy_plane1_test_rod120"),
    ("trajectory_circle_optical_table", "legacy_plane1_test_rod120"),
    ("trajectory_circle_desk", "legacy_plane1_rod217"),
    ("trajectory_square_optical_table", "legacy_plane1_test_rod120"),
    ("trajectory_square_desk", "legacy_plane1_rod217"),
    ("3comb_optical_table", "legacy_plane1_test_rod120"),
    ("comb_optical_table", "legacy_plane1_test_rod120"),
    ("comb_desk", "legacy_plane1_rod217"),
    ("sphere_optical_table", "legacy_plane1_test_rod120"),
    ("trajectory_sphere_desk", "legacy_plane1_rod217"),
    ("inner_sphere", "legacy_plane1_test_rod120"),
    ("outer_sphere", "legacy_plane1_test_rod120"),
    ("trajectory_snake3d_optical_table_part", "legacy_plane1_test_p10_rod120"),
    ("snake_quat_fixedroll", "legacy_plane1_p16_rod194"),
    ("trajectory_snake3d_desk", "legacy_plane1_rod217"),
    ("trajectory_snake3d_optical_table", "legacy_plane1_test_p10_rod120"),
    ("snake_quat_fixedroll_2", "legacy_plane1_p16_rod194"),
)
TRAJECTORY_ASSIGNMENTS = dict(_TRAJECTORY_PROFILE_ASSIGNMENTS)
if len(TRAJECTORY_ASSIGNMENTS) != len(_TRAJECTORY_PROFILE_ASSIGNMENTS):
    raise RuntimeError("legacy trajectory profile assignment contains a duplicate ID")
if set(EXPECTED_SOURCE_SHA256) != {
    f"{trajectory_id}.xlsx" for trajectory_id in TRAJECTORY_ASSIGNMENTS
}:
    raise RuntimeError(
        "audited XLSX SHA-256 registry must exactly match the trajectory assignment"
    )


@dataclass(frozen=True)
class ImportedTrajectory:
    source_name: str
    profile_id: str
    destination_environment: str
    destination_name: str
    rod_length_m: float
    plane_pose_id: str
    source_sha256: str
    output_sha256: str
    row_count: int
    duration_ns: int
    hold_row_count: int
    orientation_source: str
    maximum_time_quantization_error_s: float


def _sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _finite_float(value: object, field: str, row_number: int) -> float:
    if isinstance(value, bool):
        raise ValueError(f"row {row_number}: {field} must be numeric")
    try:
        parsed = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row_number}: {field} must be numeric") from error
    if not math.isfinite(parsed):
        raise ValueError(f"row {row_number}: {field} must be finite")
    return parsed


def _time_to_ns(value: object, row_number: int) -> tuple[int, float]:
    time_s = _finite_float(value, "Time(s)", row_number)
    if time_s < 0.0:
        raise ValueError(f"row {row_number}: Time(s) must be non-negative")
    exact_ns = Decimal.from_float(time_s) * Decimal(1_000_000_000)
    rounded_ns = int(exact_ns.to_integral_value(rounding=ROUND_HALF_EVEN))
    error_s = float(abs(exact_ns - Decimal(rounded_ns)) / Decimal(1_000_000_000))
    if error_s > 0.500001e-9:
        raise ValueError(
            f"row {row_number}: Time(s) is more than 0.5 ns from ROS Duration"
        )
    return rounded_ns, error_s


def _z_vector_to_quaternion(values: Iterable[float], row_number: int) -> tuple[float, ...]:
    z = np.asarray(tuple(values), dtype=float)
    norm = float(np.linalg.norm(z))
    if not math.isfinite(norm) or norm < 1.0e-9:
        raise ValueError(f"row {row_number}: Z-axis direction is zero or non-finite")
    z_axis = z / norm
    up = np.asarray((0.0, 1.0, 0.0), dtype=float)
    if abs(float(np.dot(z_axis, up))) > 0.999:
        up = np.asarray((1.0, 0.0, 0.0), dtype=float)
    x_axis = np.cross(up, z_axis)
    x_axis /= np.linalg.norm(x_axis)
    y_axis = np.cross(z_axis, x_axis)
    quaternion = Rotation.from_matrix(np.column_stack((x_axis, y_axis, z_axis))).as_quat()
    return tuple(float(value) for value in quaternion)


def _source_rows(
    path: Path,
    source_payload: bytes | None = None,
) -> tuple[list[dict[str, object]], str]:
    source = io.BytesIO(source_payload) if source_payload is not None else path
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError(f"{path.name}: expected exactly one worksheet")
        worksheet = workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=False)
        try:
            header_cells = next(iterator)
        except StopIteration as error:
            raise ValueError(f"{path.name}: workbook is empty") from error
        headers = [cell.value for cell in header_cells]
        if any(not isinstance(value, str) or not value for value in headers):
            raise ValueError(f"{path.name}: header cells must be non-empty strings")
        if len(set(headers)) != len(headers):
            raise ValueError(f"{path.name}: duplicate header")
        required = {"Seq", "Time(s)", "x", "y", "z"}
        quaternion_columns = {"qx", "qy", "qz", "qw"}
        z_vector_columns = {"Zbx", "Zby", "Zbz"}
        available = set(headers)
        if not required.issubset(available):
            raise ValueError(f"{path.name}: missing required columns")
        if quaternion_columns.issubset(available):
            orientation_source = "quaternion_xyzw"
        elif z_vector_columns.issubset(available):
            orientation_source = "body_z_vector"
        else:
            raise ValueError(f"{path.name}: missing supported orientation columns")
        rows: list[dict[str, object]] = []
        for excel_row_number, cells in enumerate(iterator, start=2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"{path.name}: formulas are not accepted")
            values = [cell.value for cell in cells]
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            rows.append(dict(zip(headers, values, strict=True)))
        if len(rows) < 2:
            raise ValueError(f"{path.name}: at least two waypoints are required")
        return rows, orientation_source
    finally:
        workbook.close()


def _render_csv(
    path: Path,
    frame_id: str,
    source_payload: bytes | None = None,
) -> tuple[bytes, dict[str, object]]:
    source_rows, orientation_source = _source_rows(path, source_payload)
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(CSV_HEADER)
    previous_pose: tuple[float, ...] | None = None
    previous_ns = -1
    maximum_quantization_error = 0.0
    hold_count = 0
    for index, source in enumerate(source_rows, start=1):
        excel_row_number = index + 1
        sequence = _finite_float(source["Seq"], "Seq", excel_row_number)
        if sequence != float(index):
            raise ValueError(
                f"{path.name} row {excel_row_number}: Seq must be exactly {index}"
            )
        time_ns, quantization_error = _time_to_ns(
            source["Time(s)"], excel_row_number
        )
        maximum_quantization_error = max(maximum_quantization_error, quantization_error)
        if index == 1 and time_ns != 0:
            raise ValueError(f"{path.name}: first Time(s) must be zero")
        if time_ns <= previous_ns:
            raise ValueError(f"{path.name}: Time(s) is not strictly increasing after ns conversion")
        previous_ns = time_ns
        position = tuple(
            _finite_float(source[field], field, excel_row_number)
            for field in ("x", "y", "z")
        )
        if orientation_source == "quaternion_xyzw":
            quaternion = tuple(
                _finite_float(source[field], field, excel_row_number)
                for field in ("qx", "qy", "qz", "qw")
            )
            norm = math.sqrt(sum(value * value for value in quaternion))
            if abs(norm - 1.0) > 1.0e-3:
                raise ValueError(
                    f"{path.name} row {excel_row_number}: quaternion norm is outside tolerance"
                )
        else:
            z_vector = tuple(
                _finite_float(source[field], field, excel_row_number)
                for field in ("Zbx", "Zby", "Zbz")
            )
            quaternion = _z_vector_to_quaternion(z_vector, excel_row_number)
        pose = (*position, *quaternion)
        phase = "hold" if previous_pose is not None and pose == previous_pose else "draw"
        if phase == "hold":
            hold_count += 1
        previous_pose = pose
        writer.writerow(
            (
                f"{time_ns // 1_000_000_000}.{time_ns % 1_000_000_000:09d}",
                frame_id,
                *(repr(value) for value in position),
                *(repr(value) for value in quaternion),
                phase,
            )
        )
    payload = output.getvalue().encode("utf-8")
    return payload, {
        "row_count": len(source_rows),
        "duration_ns": previous_ns,
        "hold_row_count": hold_count,
        "orientation_source": orientation_source,
        "maximum_time_quantization_error_s": maximum_quantization_error,
    }


def _write_once(path: Path, payload: bytes) -> None:
    if path.exists():
        if path.read_bytes() != payload:
            raise FileExistsError(
                f"refusing to overwrite authoritative file with different bytes: {path}"
            )
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def _manifest_payload(
    environment_id: str,
    imports: list[ImportedTrajectory],
    source_provenance: Mapping[str, str],
) -> bytes:
    profile = TEST_PROFILES[environment_id]
    plane_pose = PLANE_POSES[profile.plane_pose_id]
    test_profile: dict[str, object] = {
        "profile_id": profile.profile_id,
        "rod_length_m": profile.rod_length_m,
        "plane_pose_id": profile.plane_pose_id,
        "plane_pose_source": (
            "user-provided UR teach-pendant base->plane1 candidate converted to "
            "base_link with the controlled URDF base_link->base Rz(pi) transform"
        ),
        "historical_per_file_frame_binding": False,
        "legacy_executor": profile.legacy_executor,
        "legacy_tcp_evidence": profile.legacy_tcp_evidence,
        "selection_basis": profile.selection_basis,
    }
    if profile.plane_pose_id == "user_measured_candidate_p16":
        test_profile.update(
            {
                "historical_per_file_frame_binding_status": "unknown_not_recovered",
                "plane_pose_selection_evidence": (
                    "p16 is the active candidate in the user-provided conversion script"
                ),
                "plane_pose_registry": {
                    "path": "environments/legacy_plane1_pose_registry.yaml",
                    "raw_base_to_plane1_pose_xyz_rotvec": list(
                        plane_pose.raw_base_to_plane1_pose_xyz_rotvec
                    ),
                    "base_link_conversion": (
                        "controlled_urdf_base_link_to_base_rz_pi"
                    ),
                    "derived_base_link_to_plane1": {
                        "translation_m": list(plane_pose.translation_m),
                        "quaternion_xyzw": list(plane_pose.quaternion_xyzw),
                        "plane_z_axis_base_link": list(
                            plane_pose.plane_z_axis_base_link
                        ),
                    },
                    "normal_direction": plane_pose.normal_direction,
                },
            }
        )
    document = {
        "schema_version": 1,
        "source": dict(source_provenance),
        "environment_id": environment_id,
        "test_profile": test_profile,
        "source_direct_admission": {
            "accepted": False,
            "reason": (
                "XLSX lacks the authoritative CSV frame/phase schema and most files "
                "contain non-integer-nanosecond time values"
            ),
        },
        "conversion_contract": {
            "derived_csv_becomes_authority": True,
            "nearest_ns_conversion_explicitly_accepted": True,
            "frame_id": "plane1",
            "preserve_row_order_and_count": True,
            "preserve_xyz_values": True,
            "time_conversion": (
                "nearest ROS integer nanosecond, maximum 0.5 ns representation error"
            ),
            "z_vector_orientation": "exact legacy vector_to_quaternion algorithm",
            "quaternion_orientation": "preserved without normalization",
            "hold_phase": "only exact consecutive Pose duplicates",
            "runtime_retiming_or_waypoint_changes": False,
        },
        "files": [
            {
                "source_file": item.source_name,
                "source_sha256": item.source_sha256,
                "output_file": item.destination_name,
                "output_sha256": item.output_sha256,
                "row_count": item.row_count,
                "duration_ns": item.duration_ns,
                "hold_row_count": item.hold_row_count,
                "orientation_source": item.orientation_source,
                "maximum_time_quantization_error_s": (
                    item.maximum_time_quantization_error_s
                ),
            }
            for item in imports
        ],
    }
    return yaml.safe_dump(document, sort_keys=False).encode("utf-8")


def import_directory(
    source_root: Path,
    workspace: Path,
    *,
    accept_nearest_ns_conversion: bool,
    trajectory_assignments: Mapping[str, str] | None = None,
    expected_source_sha256: Mapping[str, str] | None = None,
    source_provenance: Mapping[str, str] | None = None,
) -> list[ImportedTrajectory]:
    if not accept_nearest_ns_conversion:
        raise ValueError(
            "XLSX times are not authoritative ROS durations; explicitly accept the "
            "documented nearest-nanosecond upstream conversion"
        )
    source_files = sorted(source_root.glob("*.xlsx"))
    if not source_files:
        raise FileNotFoundError(f"no XLSX files found under {source_root}")
    if trajectory_assignments is None:
        if expected_source_sha256 is not None or source_provenance is not None:
            raise ValueError(
                "the fixed source import does not permit provenance or SHA overrides"
            )
        assignments = dict(TRAJECTORY_ASSIGNMENTS)
        source_hash_registry = dict(EXPECTED_SOURCE_SHA256)
        manifest_source_provenance = dict(SOURCE_PROVENANCE)
    else:
        if expected_source_sha256 is None or source_provenance is None:
            raise ValueError(
                "custom trajectory assignments require explicit source SHA-256 "
                "and non-fixed source provenance"
            )
        assignments = dict(trajectory_assignments)
        source_hash_registry = dict(expected_source_sha256)
        manifest_source_provenance = dict(source_provenance)
        if manifest_source_provenance == SOURCE_PROVENANCE:
            raise ValueError(
                "custom trajectory assignments cannot claim the fixed audited commit"
            )
        if set(manifest_source_provenance) != set(SOURCE_PROVENANCE) or any(
            not isinstance(value, str) or not value
            for value in manifest_source_provenance.values()
        ):
            raise ValueError(
                "custom source provenance must contain non-empty repository, branch, "
                "commit, and path strings"
            )
    source_names = {path.stem for path in source_files}
    assignment_names = set(assignments)
    if source_names != assignment_names:
        missing = sorted(assignment_names - source_names)
        unexpected = sorted(source_names - assignment_names)
        raise ValueError(
            "source XLSX set does not exactly match the audited trajectory assignment; "
            f"missing={missing}, unexpected={unexpected}"
        )
    expected_source_names = {f"{name}.xlsx" for name in assignments}
    if set(source_hash_registry) != expected_source_names:
        missing = sorted(expected_source_names - set(source_hash_registry))
        unexpected = sorted(set(source_hash_registry) - expected_source_names)
        raise ValueError(
            "source SHA-256 registry does not exactly match the selected XLSX set; "
            f"missing={missing}, unexpected={unexpected}"
        )
    for source_name, digest in source_hash_registry.items():
        if (
            not isinstance(digest, str)
            or len(digest) != 64
            or any(character not in "0123456789abcdef" for character in digest)
        ):
            raise ValueError(f"{source_name}: expected SHA-256 must be lowercase hex")
        fixed_digest = EXPECTED_SOURCE_SHA256.get(source_name)
        if fixed_digest is not None and digest != fixed_digest:
            raise ValueError(
                f"{source_name}: the audited filename cannot override its fixed "
                f"SHA-256 for {SOURCE_COMMIT}"
            )
    unknown_environments = sorted(set(assignments.values()) - set(TEST_PROFILES))
    if unknown_environments:
        raise ValueError(
            "trajectory assignment references unknown test profiles: "
            + ", ".join(unknown_environments)
        )

    # Snapshot and authenticate every audited workbook before parsing or writing
    # anything.  Parsing from the same bytes closes the check/use race and prevents
    # arbitrary same-named XLSX files from being attributed to SOURCE_COMMIT.
    source_payloads: dict[str, bytes] = {}
    source_hashes: dict[str, str] = {}
    for source_path in source_files:
        payload = source_path.read_bytes()
        actual_sha256 = _sha256_bytes(payload)
        expected_sha256 = source_hash_registry[source_path.name]
        if actual_sha256 != expected_sha256:
            raise ValueError(
                f"{source_path.name}: source SHA-256 mismatch; expected "
                f"{expected_sha256}, got {actual_sha256}"
            )
        source_payloads[source_path.name] = payload
        source_hashes[source_path.name] = actual_sha256

    imported: list[ImportedTrajectory] = []
    for source_path in source_files:
        payload, metadata = _render_csv(
            source_path,
            "plane1",
            source_payloads[source_path.name],
        )
        environment_id = assignments[source_path.stem]
        profile = TEST_PROFILES[environment_id]
        destination = (
            workspace
            / "environments"
            / environment_id
            / "trajectories"
            / f"{source_path.stem}.csv"
        )
        _write_once(destination, payload)
        imported.append(
            ImportedTrajectory(
                source_name=source_path.name,
                profile_id=profile.profile_id,
                destination_environment=environment_id,
                destination_name=destination.name,
                rod_length_m=profile.rod_length_m,
                plane_pose_id=profile.plane_pose_id,
                source_sha256=source_hashes[source_path.name],
                output_sha256=_sha256_bytes(payload),
                row_count=int(metadata["row_count"]),
                duration_ns=int(metadata["duration_ns"]),
                hold_row_count=int(metadata["hold_row_count"]),
                orientation_source=str(metadata["orientation_source"]),
                maximum_time_quantization_error_s=float(
                    metadata["maximum_time_quantization_error_s"]
                ),
            )
        )
    for environment_id in sorted({item.destination_environment for item in imported}):
        environment_imports = [
            item for item in imported if item.destination_environment == environment_id
        ]
        manifest_path = (
            workspace
            / "environments"
            / environment_id
            / "trajectories"
            / "localizaiton_ros2_import.yaml"
        )
        _write_once(
            manifest_path,
            _manifest_payload(
                environment_id,
                environment_imports,
                manifest_source_provenance,
            ),
        )
    return imported


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument(
        "--accept-nearest-ns-conversion",
        action="store_true",
        required=True,
        help=(
            "explicitly create new authority times at the nearest ROS nanosecond; "
            "this is an upstream conversion, never a runtime repair"
        ),
    )
    args = parser.parse_args()
    imports = import_directory(
        args.source_root.resolve(),
        args.workspace.resolve(),
        accept_nearest_ns_conversion=args.accept_nearest_ns_conversion,
    )
    for item in imports:
        print(
            f"{item.source_name} -> {item.profile_id}/"
            f"{item.destination_environment}/"
            f"trajectories/{item.destination_name} rows={item.row_count} "
            f"rod_length_m={item.rod_length_m:.3f} plane_pose_id={item.plane_pose_id} "
            f"duration_s={item.duration_ns / 1.0e9:.9f} sha256={item.output_sha256}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
