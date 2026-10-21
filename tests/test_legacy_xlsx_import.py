from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import shutil
import subprocess
import sys
from pathlib import Path
from xml.etree import ElementTree

import numpy as np
import pytest
import yaml
from openpyxl import Workbook, load_workbook
from scipy.spatial.transform import Rotation

from hx_waypoint_curobo.environment_builder import EnvironmentBuilder
from hx_waypoint_curobo.robot_tool_builder import RobotToolBuilder
from hx_waypoint_curobo.trajectory_io import read_timed_pose_csv
from scripts.import_localizaiton_ros2_data import (
    EXPECTED_SOURCE_SHA256,
    PLANE_POSES,
    PLANE_POSE_REGISTRY_PATH,
    SOURCE_BRANCH,
    SOURCE_COMMIT,
    SOURCE_PATH,
    SOURCE_REPOSITORY,
    TEST_PROFILES,
    TRAJECTORY_ASSIGNMENTS,
    import_directory,
)


PROFILE_TRAJECTORIES = {
    "legacy_plane1_rod217": {
        "8omega3_desk",
        "comb_desk",
        "omega3_desk",
        "trajectory_circle_desk",
        "trajectory_snake3d_desk",
        "trajectory_sphere_desk",
        "trajectory_square_desk",
    },
    "legacy_plane1_test_rod120": {
        "3comb_optical_table",
        "comb_optical_table",
        "inner_sphere",
        "omega1",
        "omega3",
        "outer_sphere",
        "sphere_optical_table",
        "trajectory_circle_optical_table",
        "trajectory_square_optical_table",
    },
    "legacy_plane1_test_p10_rod120": {
        "trajectory_snake3d_optical_table",
        "trajectory_snake3d_optical_table_part",
    },
    "legacy_plane1_p16_rod194": {
        "snake_quat_fixedroll",
        "snake_quat_fixedroll_2",
    },
}

PROFILE_EXPECTATIONS = {
    "legacy_plane1_rod217": {
        "profile_id": "p15_rod217_desk",
        "pose_id": "user_measured_candidate_p15",
        "rod_length_m": 0.217,
        "raw_pose": (-0.00327, -0.50126, 0.27942, -0.00576, -0.00076, 0.7778),
        "sphere_count": 20,
    },
    "legacy_plane1_test_rod120": {
        "profile_id": "p15_rod120_optical",
        "pose_id": "user_measured_candidate_p15",
        "rod_length_m": 0.120,
        "raw_pose": (-0.00327, -0.50126, 0.27942, -0.00576, -0.00076, 0.7778),
        "sphere_count": 16,
    },
    "legacy_plane1_test_p10_rod120": {
        "profile_id": "p10_rod120_large_optical",
        "pose_id": "user_measured_candidate_p10",
        "rod_length_m": 0.120,
        "raw_pose": (-0.07584, -0.38257, 0.28285, -0.00534, 0.00519, 0.82389),
        "sphere_count": 16,
    },
    "legacy_plane1_p16_rod194": {
        "profile_id": "p16_rod194_quaternion_flipped_normal",
        "pose_id": "user_measured_candidate_p16",
        "rod_length_m": 0.194,
        "raw_pose": (-0.03413, -0.38168, 0.33594, -1.25152, 2.88032, -0.01303),
        "sphere_count": 18,
    },
}

RAW_USER_POSES = (
    (-0.00231, -0.53681, 0.28718, 0.00568, 0.00732, 0.82367),
    (-0.05497, -0.42631, 0.20931, -0.0014, 0.0117, 0.81032),
    (-0.13054, -0.48361, 0.28742, -0.01593, 0.00207, 0.72029),
    (0.43, -0.98712, 0.28679, -0.01416, 0.01658, 0.81215),
    (0.43, -0.98712, 0.28679, 0.06765, -0.02718, 0.80704),
    (-0.10879, -0.55163, 0.3754, 0.00687, -0.00884, 0.80169),
    (-0.22399, -0.48041, 0.37757, -0.01239, 0.00612, 0.7976),
    (0.43942, -0.98923, 0.27947, -0.01233, -0.0007, 0.79383),
    (-0.11414, -0.51115, 0.34236, -0.00267, -0.00445, 0.83037),
    (-0.07584, -0.38257, 0.28285, -0.00534, 0.00519, 0.82389),
    (0.03841, -0.49409, 0.43719, 0.00899, -0.00547, 0.82071),
    (-0.06635, -0.38358, 0.43729, -0.0045, 0.00384, 0.84776),
    (-0.06309, -0.39567, 0.28901, -0.01561, -0.00214, 0.84671),
    (0.43942, -0.98923, 0.27947, -0.01233, -0.0007, 0.79383),
    (-0.00327, -0.50126, 0.27942, -0.00576, -0.00076, 0.7778),
    (-0.03413, -0.38168, 0.33594, -1.25152, 2.88032, -0.01303),
)


def _write_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _flatten_assignments() -> dict[str, str]:
    return {
        trajectory_id: environment_id
        for environment_id, trajectory_ids in PROFILE_TRAJECTORIES.items()
        for trajectory_id in trajectory_ids
    }


SYNTHETIC_SOURCE_PROVENANCE = {
    "repository": "pytest-generated-fixture",
    "branch": "synthetic",
    "commit": "not-a-git-commit",
    "path": "generated-workbooks",
}


def _synthetic_import_kwargs(*paths: Path) -> dict[str, object]:
    return {
        "expected_source_sha256": {path.name: _sha256(path) for path in paths},
        "source_provenance": SYNTHETIC_SOURCE_PROVENANCE,
    }


def _copy_matrix_validation_workspace(
    repository_root: Path,
    destination: Path,
) -> Path:
    workspace = destination / "matrix-workspace"
    (workspace / "scripts").mkdir(parents=True)
    (workspace / "environments").mkdir()
    for script_name in (
        "import_localizaiton_ros2_data.py",
        "run_localizaiton_ros2_compile_matrix.sh",
    ):
        shutil.copy2(
            repository_root / "scripts" / script_name,
            workspace / "scripts" / script_name,
        )
    shutil.copy2(
        PLANE_POSE_REGISTRY_PATH,
        workspace / "environments" / PLANE_POSE_REGISTRY_PATH.name,
    )
    for environment_id in PROFILE_TRAJECTORIES:
        source_environment = repository_root / "environments" / environment_id
        destination_environment = workspace / "environments" / environment_id
        destination_environment.mkdir()
        for config_name in ("environment.yaml", "setup.yaml"):
            shutil.copy2(
                source_environment / config_name,
                destination_environment / config_name,
            )
        shutil.copytree(
            source_environment / "trajectories",
            destination_environment / "trajectories",
        )
    python_dir = workspace / ".venv" / "bin"
    python_dir.mkdir(parents=True)
    (python_dir / "python").symlink_to(Path(sys.executable).resolve())
    return workspace


def _run_matrix_validation(workspace: Path) -> subprocess.CompletedProcess[str]:
    environment = dict(os.environ)
    environment["UR_WAYPOINT_ROOT"] = str(workspace)
    return subprocess.run(
        [
            str(workspace / "scripts" / "run_localizaiton_ros2_compile_matrix.sh"),
            "--validate-only",
        ],
        cwd=workspace,
        env=environment,
        check=False,
        capture_output=True,
        text=True,
    )


def _run_matrix_full(
    workspace: Path,
    output_path: Path,
) -> subprocess.CompletedProcess[str]:
    environment = dict(os.environ)
    environment["UR_WAYPOINT_ROOT"] = str(workspace)
    return subprocess.run(
        [
            str(workspace / "scripts" / "run_localizaiton_ros2_compile_matrix.sh"),
            str(output_path),
        ],
        cwd=workspace,
        env=environment,
        check=False,
        capture_output=True,
        text=True,
    )


def test_import_preserves_rows_and_refuses_authority_overwrite(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    workspace = tmp_path / "workspace"
    z_environment = "legacy_plane1_test_rod120"
    q_environment = "legacy_plane1_p16_rod194"
    for environment_id in (z_environment, q_environment):
        (workspace / "environments" / environment_id / "trajectories").mkdir(
            parents=True
        )

    z_vector_source = source_root / "z_path.xlsx"
    _write_workbook(
        z_vector_source,
        ["Seq", "Time(s)", "x", "y", "z", "Zbx", "Zby", "Zbz"],
        [
            [1, 0.0, 0.0, 0.0, 0.06, 0.0, 0.0, -1.0],
            [2, 0.01, 0.01, 0.0, 0.06, 0.0, 0.0, -1.0],
            [3, 0.02, 0.01, 0.0, 0.06, 0.0, 0.0, -1.0],
        ],
    )
    quaternion_source = source_root / "q_path.xlsx"
    _write_workbook(
        quaternion_source,
        ["Seq", "Time(s)", "x", "y", "z", "qx", "qy", "qz", "qw"],
        [
            [1, 0.0, -0.05, -0.05, -0.05, 1.0, 0.0, 0.0, 0.0],
            [2, 1.0 / 30.0, -0.04, -0.05, -0.05, 1.0, 0.0, 0.0, 0.0],
        ],
    )
    assignments = {"z_path": z_environment, "q_path": q_environment}
    synthetic_kwargs = _synthetic_import_kwargs(
        z_vector_source,
        quaternion_source,
    )

    with pytest.raises(ValueError, match="explicitly accept"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=False,
            trajectory_assignments=assignments,
            **synthetic_kwargs,
        )
    with pytest.raises(ValueError, match="does not exactly match"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
        )
    with pytest.raises(ValueError, match="custom trajectory assignments require"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
            trajectory_assignments=assignments,
        )
    imported = import_directory(
        source_root,
        workspace,
        accept_nearest_ns_conversion=True,
        trajectory_assignments=assignments,
        **synthetic_kwargs,
    )
    assert {item.destination_environment for item in imported} == {
        z_environment,
        q_environment,
    }

    z_csv = workspace / "environments" / z_environment / "trajectories" / "z_path.csv"
    with z_csv.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 3
    assert [row["time_s"] for row in rows] == [
        "0.000000000",
        "0.010000000",
        "0.020000000",
    ]
    assert rows[2]["phase"] == "hold"
    assert rows[2]["tip_x_m"] == rows[1]["tip_x_m"]
    assert rows[0]["tip_qy"] == "1.0"

    q_csv = workspace / "environments" / q_environment / "trajectories" / "q_path.csv"
    with q_csv.open(newline="", encoding="utf-8") as stream:
        q_rows = list(csv.DictReader(stream))
    assert q_rows[0]["tip_qx"] == "1.0"
    assert q_rows[0]["tip_qw"] == "0.0"
    assert q_rows[1]["time_s"] == "0.033333333"

    manifest_paths = {
        environment_id: workspace
        / "environments"
        / environment_id
        / "trajectories"
        / "localizaiton_ros2_import.yaml"
        for environment_id in (z_environment, q_environment)
    }
    z_manifest = yaml.safe_load(manifest_paths[z_environment].read_text(encoding="utf-8"))
    q_manifest = yaml.safe_load(manifest_paths[q_environment].read_text(encoding="utf-8"))
    assert z_manifest["test_profile"]["profile_id"] == "p15_rod120_optical"
    assert z_manifest["conversion_contract"]["preserve_row_order_and_count"] is True
    assert (
        z_manifest["conversion_contract"]["runtime_retiming_or_waypoint_changes"]
        is False
    )
    assert q_manifest["test_profile"]["profile_id"] == (
        "p16_rod194_quaternion_flipped_normal"
    )
    assert q_manifest["test_profile"]["historical_per_file_frame_binding_status"] == (
        "unknown_not_recovered"
    )
    assert q_manifest["test_profile"]["plane_pose_registry"]["normal_direction"] == (
        "base_link_negative_z"
    )
    assert z_manifest["source"] == SYNTHETIC_SOURCE_PROVENANCE
    assert q_manifest["source"] == SYNTHETIC_SOURCE_PROVENANCE

    generated_paths = (z_csv, q_csv, *manifest_paths.values())
    before = {
        path: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in generated_paths
    }
    import_directory(
        source_root,
        workspace,
        accept_nearest_ns_conversion=True,
        trajectory_assignments=assignments,
        **synthetic_kwargs,
    )
    after = {
        path: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in generated_paths
    }
    assert after == before

    z_csv_bytes = z_csv.read_bytes()
    z_csv.write_bytes(z_csv_bytes + b"tampered-authority\n")
    with pytest.raises(FileExistsError, match="z_path.csv"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
            trajectory_assignments=assignments,
            **synthetic_kwargs,
        )
    z_csv.write_bytes(z_csv_bytes)

    quaternion_source_bytes = quaternion_source.read_bytes()
    workbook = load_workbook(quaternion_source)
    workbook.properties.title = "source bytes changed without changing trajectory rows"
    workbook.save(quaternion_source)
    workbook.close()
    with pytest.raises(ValueError, match="q_path.xlsx: source SHA-256 mismatch"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
            trajectory_assignments=assignments,
            **synthetic_kwargs,
        )
    quaternion_source.write_bytes(quaternion_source_bytes)

    workbook = load_workbook(z_vector_source)
    workbook.active["C2"] = 0.5
    workbook.save(z_vector_source)
    workbook.close()
    with pytest.raises(ValueError, match="z_path.xlsx: source SHA-256 mismatch"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
            trajectory_assignments=assignments,
            **synthetic_kwargs,
        )


def test_cli_requires_explicit_nearest_ns_acceptance(
    tmp_path: Path,
    repository_root: Path,
) -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(repository_root / "scripts" / "import_localizaiton_ros2_data.py"),
            "--source-root",
            str(tmp_path),
            "--workspace",
            str(tmp_path),
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 2
    assert "--accept-nearest-ns-conversion" in result.stderr


def test_profile_assignment_is_exactly_twenty_unique_trajectories() -> None:
    expected = _flatten_assignments()
    assert len(expected) == 20
    assert expected == TRAJECTORY_ASSIGNMENTS
    assert set(TEST_PROFILES) == set(PROFILE_TRAJECTORIES)
    assert {profile.rod_length_m for profile in TEST_PROFILES.values()} == {
        0.120,
        0.194,
        0.217,
    }
    assert all("068" not in profile.profile_id for profile in TEST_PROFILES.values())
    assert "legacy_plane1_rod194" not in TEST_PROFILES
    assert {
        TRAJECTORY_ASSIGNMENTS["snake_quat_fixedroll"],
        TRAJECTORY_ASSIGNMENTS["snake_quat_fixedroll_2"],
    } == {"legacy_plane1_p16_rod194"}


def test_fixed_commit_source_hashes_match_every_active_import_manifest(
    repository_root: Path,
) -> None:
    assert len(EXPECTED_SOURCE_SHA256) == 20
    assert set(EXPECTED_SOURCE_SHA256) == {
        f"{trajectory_id}.xlsx" for trajectory_id in TRAJECTORY_ASSIGNMENTS
    }
    manifested_source_hashes: dict[str, str] = {}
    expected_provenance = {
        "repository": SOURCE_REPOSITORY,
        "branch": SOURCE_BRANCH,
        "commit": SOURCE_COMMIT,
        "path": SOURCE_PATH,
    }
    for environment_id in PROFILE_TRAJECTORIES:
        manifest_path = (
            repository_root
            / "environments"
            / environment_id
            / "trajectories"
            / "localizaiton_ros2_import.yaml"
        )
        manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
        assert manifest["source"] == expected_provenance
        for entry in manifest["files"]:
            source_file = entry["source_file"]
            assert source_file not in manifested_source_hashes
            manifested_source_hashes[source_file] = entry["source_sha256"]
    assert manifested_source_hashes == EXPECTED_SOURCE_SHA256


def test_fixed_commit_import_rejects_same_named_tampering_before_any_write(
    tmp_path: Path,
) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    for source_name in EXPECTED_SOURCE_SHA256:
        (source_root / source_name).write_bytes(b"tampered XLSX bytes")
    workspace = tmp_path / "workspace"

    with pytest.raises(ValueError, match="source SHA-256 mismatch"):
        import_directory(
            source_root,
            workspace,
            accept_nearest_ns_conversion=True,
        )

    assert not workspace.exists()


def test_pose_registry_preserves_all_user_rows_and_derives_rz_pi() -> None:
    document = yaml.safe_load(PLANE_POSE_REGISTRY_PATH.read_text(encoding="utf-8"))
    assert document["source"]["supplied_row_count"] == 16
    assert document["base_link_conversion"]["conversion_id"] == (
        "controlled_urdf_base_link_to_base_rz_pi"
    )
    assert len(PLANE_POSES) == len(RAW_USER_POSES) == 16

    base_link_to_base = Rotation.from_euler("z", math.pi)
    for source_row, expected_raw_pose in enumerate(RAW_USER_POSES, start=1):
        pose_id = f"user_measured_candidate_p{source_row}"
        pose = PLANE_POSES[pose_id]
        registry_entry = document["poses"][pose_id]
        assert pose.source_row == source_row
        assert pose.raw_base_to_plane1_pose_xyz_rotvec == expected_raw_pose
        assert registry_entry["raw_base_to_plane1_pose_xyz_rotvec"] == list(
            expected_raw_pose
        )

        expected_translation = base_link_to_base.apply(expected_raw_pose[:3])
        expected_rotation = base_link_to_base * Rotation.from_rotvec(
            expected_raw_pose[3:]
        )
        assert pose.translation_m == pytest.approx(expected_translation, abs=1.0e-12)
        assert abs(
            float(np.dot(pose.quaternion_xyzw, expected_rotation.as_quat()))
        ) == pytest.approx(1.0, abs=1.0e-12)
        assert pose.plane_z_axis_base_link == pytest.approx(
            expected_rotation.apply((0.0, 0.0, 1.0)),
            abs=1.0e-12,
        )
        assert pose.quaternion_xyzw[3] >= 0.0

    p8 = PLANE_POSES["user_measured_candidate_p8"]
    p14 = PLANE_POSES["user_measured_candidate_p14"]
    assert p14.exact_duplicate_of == p8.pose_id
    assert p14.raw_base_to_plane1_pose_xyz_rotvec == p8.raw_base_to_plane1_pose_xyz_rotvec
    assert p14.translation_m == p8.translation_m
    assert p14.quaternion_xyzw == p8.quaternion_xyzw

    assert all(
        PLANE_POSES[f"user_measured_candidate_p{index}"].plane_z_axis_base_link[2]
        > 0.99
        for index in range(1, 16)
    )
    p16 = PLANE_POSES["user_measured_candidate_p16"]
    assert p16.normal_direction == "base_link_negative_z"
    assert p16.plane_z_axis_base_link[2] < -0.999


def test_measured_profile_environments_load_and_are_immutable(
    repository_root: Path,
) -> None:
    all_trajectory_ids: set[str] = set()
    immutable_paths: list[Path] = []

    for environment_id, expected_trajectories in PROFILE_TRAJECTORIES.items():
        root = repository_root / "environments" / environment_id
        csv_paths = sorted((root / "trajectories").glob("*.csv"))
        assert {path.stem for path in csv_paths} == expected_trajectories
        assert not (all_trajectory_ids & expected_trajectories)
        all_trajectory_ids.update(expected_trajectories)

        manifest_path = root / "trajectories" / "localizaiton_ros2_import.yaml"
        immutable_paths.extend((*csv_paths, manifest_path))
        manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
        expected = PROFILE_EXPECTATIONS[environment_id]
        assert manifest["environment_id"] == environment_id
        assert manifest["test_profile"]["profile_id"] == expected["profile_id"]
        assert manifest["test_profile"]["rod_length_m"] == expected["rod_length_m"]
        assert manifest["test_profile"]["plane_pose_id"] == expected["pose_id"]
        assert manifest["test_profile"]["historical_per_file_frame_binding"] is False
        if environment_id == "legacy_plane1_p16_rod194":
            assert manifest["test_profile"][
                "historical_per_file_frame_binding_status"
            ] == "unknown_not_recovered"
            assert manifest["test_profile"]["plane_pose_selection_evidence"] == (
                "p16 is the active candidate in the user-provided conversion script"
            )
            assert manifest["test_profile"]["plane_pose_registry"][
                "normal_direction"
            ] == "base_link_negative_z"
        assert {item["output_file"] for item in manifest["files"]} == {
            f"{trajectory_id}.csv" for trajectory_id in expected_trajectories
        }
        for item in manifest["files"]:
            assert _sha256(root / "trajectories" / item["output_file"]) == item[
                "output_sha256"
            ]

    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in immutable_paths}
    for environment_id in PROFILE_TRAJECTORIES:
        environment = EnvironmentBuilder().load(repository_root / "environments" / environment_id)
        expected = PROFILE_EXPECTATIONS[environment_id]
        profile = environment.environment["legacy_test_profile"]
        assert environment.environment_id == environment_id
        assert profile["profile_id"] == expected["profile_id"]
        assert profile["historical_per_file_frame_binding"] is False
        assert profile["base_link_conversion"] == (
            "controlled_urdf_base_link_to_base_rz_pi"
        )
        registered_pose = PLANE_POSES[expected["pose_id"]]
        assert profile["raw_base_to_plane1_pose_xyz_rotvec"] == pytest.approx(
            registered_pose.raw_base_to_plane1_pose_xyz_rotvec
        )
        plane = environment.frames["frames"]["plane1"]
        assert plane["translation_m"] == pytest.approx(
            registered_pose.translation_m,
            abs=1.0e-12,
        )
        assert abs(
            float(
                np.dot(
                    np.asarray(plane["quaternion_xyzw"], dtype=float),
                    np.asarray(registered_pose.quaternion_xyzw, dtype=float),
                )
            )
        ) == pytest.approx(1.0, abs=1.0e-12)
        if environment_id == "legacy_plane1_p16_rod194":
            assert profile["plane_normal_direction"] == "base_link_negative_z"
            assert profile["plane_pose_registry"] == (
                "environments/legacy_plane1_pose_registry.yaml"
            )
        assert environment.tcp["tool"]["translation_m"] == [
            0.0,
            0.0,
            expected["rod_length_m"],
        ]
        assert environment.tcp["rod"]["length_m"] == expected["rod_length_m"]
        assert environment.zones["contact_policy"]["allowed_surface"] == (
            "safety_lower_plane"
        )
        assert environment.zones["analysis_zones"]["safety_lower_plane"][
            "offset_m"
        ] == -0.25

        spheres = RobotToolBuilder(environment).collision_spheres()
        assert len(spheres) == expected["sphere_count"]
        assert all(sphere.link_name == "tool0" for sphere in spheres)
        assert all(
            sphere.radius_m == environment.tcp["rod"]["radius_m"] for sphere in spheres
        )
        centers = np.asarray([sphere.center_m[2] for sphere in spheres])
        spacing = float(environment.tcp["rod"]["sphere_spacing_m"])
        assert np.all(np.diff(centers) <= spacing + 1.0e-12)
        assert centers[-1] + spheres[-1].radius_m == pytest.approx(
            expected["rod_length_m"],
            abs=1.0e-12,
        )

        for trajectory_id in PROFILE_TRAJECTORIES[environment_id]:
            poses = read_timed_pose_csv(
                environment.root / "trajectories" / f"{trajectory_id}.csv"
            )
            assert len(poses) >= 2
            assert poses[0].time_s == 0.0
            assert {pose.frame_id for pose in poses} == {"plane1"}
            assert min(pose.position_m[2] for pose in poses) > -0.25

    after = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in immutable_paths}
    assert after == before
    assert all_trajectory_ids == set(TRAJECTORY_ASSIGNMENTS)


def test_p10_rod194_environment_remains_historical_evidence_only(
    repository_root: Path,
) -> None:
    historical_root = repository_root / "environments" / "legacy_plane1_rod194"
    current_root = repository_root / "environments" / "legacy_plane1_p16_rod194"
    historical = EnvironmentBuilder().load(historical_root)
    profile = historical.environment["legacy_test_profile"]
    assert profile["profile_id"] == "p10_rod194_quaternion"
    assert profile["plane_pose_id"] == "user_measured_candidate_p10"
    assert historical.environment_id == "legacy_plane1_rod194"
    assert "legacy_plane1_rod194" not in TEST_PROFILES
    assert "legacy_plane1_rod194" not in set(TRAJECTORY_ASSIGNMENTS.values())

    for trajectory_id in ("snake_quat_fixedroll", "snake_quat_fixedroll_2"):
        historical_csv = historical_root / "trajectories" / f"{trajectory_id}.csv"
        current_csv = current_root / "trajectories" / f"{trajectory_id}.csv"
        assert _sha256(historical_csv) == _sha256(current_csv)
        assert historical_csv.read_bytes() == current_csv.read_bytes()


def test_teach_pendant_base_pose_is_converted_through_controlled_urdf(
    repository_root: Path,
) -> None:
    urdf_path = (
        repository_root
        / "environments"
        / "legacy_plane1_test_rod120"
        / "robot"
        / "assets"
        / "ur5.urdf"
    )
    urdf = ElementTree.parse(urdf_path).getroot()
    joint = urdf.find("./joint[@name='base_link-base_fixed_joint']")
    assert joint is not None
    assert joint.find("parent").get("link") == "base_link"
    assert joint.find("child").get("link") == "base"
    origin = joint.find("origin")
    xyz = np.asarray([float(value) for value in origin.get("xyz").split()])
    rpy = np.asarray([float(value) for value in origin.get("rpy").split()])
    assert xyz == pytest.approx((0.0, 0.0, 0.0), abs=1.0e-12)
    assert rpy == pytest.approx((0.0, 0.0, math.pi), abs=1.0e-12)
    base_link_to_base = Rotation.from_euler("xyz", rpy)

    for environment_id, expected in PROFILE_EXPECTATIONS.items():
        environment = EnvironmentBuilder().load(
            repository_root / "environments" / environment_id
        )
        raw_pose = np.asarray(expected["raw_pose"], dtype=float)
        registered_pose = PLANE_POSES[expected["pose_id"]]
        assert raw_pose == pytest.approx(
            registered_pose.raw_base_to_plane1_pose_xyz_rotvec
        )
        profile = environment.environment["legacy_test_profile"]
        assert profile["raw_base_to_plane1_pose_xyz_rotvec"] == pytest.approx(raw_pose)
        expected_translation = base_link_to_base.apply(raw_pose[:3])
        expected_rotation = base_link_to_base * Rotation.from_rotvec(raw_pose[3:])
        plane = environment.frames["frames"]["plane1"]
        actual_translation = np.asarray(plane["translation_m"], dtype=float)
        actual_quaternion = np.asarray(plane["quaternion_xyzw"], dtype=float)
        assert actual_translation == pytest.approx(expected_translation, abs=1.0e-12)
        assert abs(float(np.dot(actual_quaternion, expected_rotation.as_quat()))) == (
            pytest.approx(1.0, abs=1.0e-12)
        )


def test_compile_matrix_validates_and_emits_actual_profile_evidence(
    repository_root: Path,
) -> None:
    script = repository_root / "scripts" / "run_localizaiton_ros2_compile_matrix.sh"
    script_text = script.read_text(encoding="utf-8")
    assert "TRAJECTORY_ASSIGNMENTS" in script_text
    assert "snake_quat_fixedroll|" not in script_text
    assert 'pipeline_status=("${PIPESTATUS[@]}")' in script_text
    assert 'if [[ "${tee_exit_code}" -ne 0 ]]' in script_text
    assert "if not torch.cuda.is_available():" in script_text
    assert "if ! nvidia-smi" in script_text
    assert "failed to initialize compile_matrix.tsv" in script_text
    assert "failed to append compile evidence" in script_text
    assert "compile_matrix.tsv.partial" in script_text
    assert '2>&1 | tee "${log_tmp}"' in script_text
    assert 'exit "${matrix_exit_code}"' in script_text
    assert "set -e" not in script_text
    environment = dict(os.environ)
    environment["UR_WAYPOINT_ROOT"] = str(repository_root)
    result = subprocess.run(
        [str(script), "--validate-only"],
        cwd=repository_root,
        env=environment,
        check=True,
        capture_output=True,
        text=True,
    )
    rows = list(csv.DictReader(io.StringIO(result.stdout), delimiter="\t"))
    assert len(rows) == 20
    assert {row["trajectory_id"] for row in rows} == set(TRAJECTORY_ASSIGNMENTS)
    assert {
        row["trajectory_id"]: row["environment_id"] for row in rows
    } == TRAJECTORY_ASSIGNMENTS

    for row in rows:
        expected = PROFILE_EXPECTATIONS[row["environment_id"]]
        assert row["profile_id"] == expected["profile_id"]
        assert float(row["rod_length_m"]) == expected["rod_length_m"]
        assert row["plane_pose_id"] == expected["pose_id"]
        registered_pose = PLANE_POSES[expected["pose_id"]]
        assert row["base_link_conversion"] == "controlled_urdf_base_link_to_base_rz_pi"
        assert json.loads(row["raw_base_to_plane1_pose_xyz_rotvec"]) == pytest.approx(
            expected["raw_pose"]
        )
        assert json.loads(row["plane_translation_m"]) == pytest.approx(
            registered_pose.translation_m,
            abs=1.0e-12,
        )
        assert json.loads(row["plane_quaternion_xyzw"]) == pytest.approx(
            registered_pose.quaternion_xyzw,
            abs=1.0e-12,
        )
        assert json.loads(row["plane_z_axis_base_link"]) == pytest.approx(
            registered_pose.plane_z_axis_base_link
        )
        assert row["plane_normal_direction"] == registered_pose.normal_direction
        source_path = (
            repository_root
            / "environments"
            / row["environment_id"]
            / "trajectories"
            / f"{row['trajectory_id']}.csv"
        )
        assert row["source_sha256"] == _sha256(source_path)
        manifest_path = source_path.parent / "localizaiton_ros2_import.yaml"
        assert row["import_manifest_sha256"] == _sha256(manifest_path)
        manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
        manifest_entry = next(
            item
            for item in manifest["files"]
            if item["output_file"] == source_path.name
        )
        assert row["source_sha256"] == manifest_entry["output_sha256"]

    quaternion_rows = {
        row["trajectory_id"]: row
        for row in rows
        if row["trajectory_id"] in {"snake_quat_fixedroll", "snake_quat_fixedroll_2"}
    }
    assert set(quaternion_rows) == {"snake_quat_fixedroll", "snake_quat_fixedroll_2"}
    assert {
        row["environment_id"] for row in quaternion_rows.values()
    } == {"legacy_plane1_p16_rod194"}
    assert {
        row["plane_normal_direction"] for row in quaternion_rows.values()
    } == {"base_link_negative_z"}


def test_compile_matrix_full_mode_rejects_existing_or_escaping_output_before_ros(
    tmp_path: Path,
    repository_root: Path,
) -> None:
    workspace = _copy_matrix_validation_workspace(repository_root, tmp_path)
    runs = workspace / "runs"
    runs.mkdir()
    existing = runs / "existing"
    existing.mkdir()
    sentinel = existing / "sentinel.txt"
    sentinel.write_text("preserve\n", encoding="utf-8")

    existing_result = _run_matrix_full(workspace, existing)

    assert existing_result.returncode == 2
    assert "evidence directory already exists" in existing_result.stderr
    assert sentinel.read_text(encoding="utf-8") == "preserve\n"

    escaped = runs / ".." / "escaped"
    escaped_result = _run_matrix_full(workspace, escaped)

    assert escaped_result.returncode == 2
    assert "new direct child" in escaped_result.stderr
    assert not (workspace / "escaped").exists()


def test_compile_matrix_rejects_csv_tampering_without_gpu(
    tmp_path: Path,
    repository_root: Path,
) -> None:
    workspace = _copy_matrix_validation_workspace(repository_root, tmp_path)
    source_path = (
        workspace
        / "environments"
        / "legacy_plane1_test_rod120"
        / "trajectories"
        / "omega3.csv"
    )
    source_path.write_bytes(source_path.read_bytes() + b"tampered\n")

    result = _run_matrix_validation(workspace)

    assert result.returncode == 2
    assert "authoritative CSV differs from import manifest" in result.stderr


@pytest.mark.parametrize(
    ("tamper", "expected_error"),
    (
        ("repository", "source provenance mismatch"),
        ("branch", "source provenance mismatch"),
        ("commit", "source provenance mismatch"),
        ("environment", "import manifest environment_id mismatch"),
        ("profile", "import manifest profile_id mismatch"),
        ("output_file", "source/output filename mismatch"),
        ("output_sha256", "authoritative CSV differs from import manifest"),
    ),
)
def test_compile_matrix_rejects_manifest_tampering_without_gpu(
    tmp_path: Path,
    repository_root: Path,
    tamper: str,
    expected_error: str,
) -> None:
    workspace = _copy_matrix_validation_workspace(repository_root, tmp_path)
    manifest_path = (
        workspace
        / "environments"
        / "legacy_plane1_test_rod120"
        / "trajectories"
        / "localizaiton_ros2_import.yaml"
    )
    manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    if tamper in {"repository", "branch", "commit"}:
        manifest["source"][tamper] = "tampered"
    elif tamper == "environment":
        manifest["environment_id"] = "tampered"
    elif tamper == "profile":
        manifest["test_profile"]["profile_id"] = "tampered"
    else:
        entry = next(
            item for item in manifest["files"] if item["output_file"] == "omega3.csv"
        )
        if tamper == "output_file":
            entry["output_file"] = "tampered.csv"
        else:
            entry["output_sha256"] = "0" * 64
    manifest_path.write_text(
        yaml.safe_dump(manifest, sort_keys=False),
        encoding="utf-8",
    )

    result = _run_matrix_validation(workspace)

    assert result.returncode == 2
    assert expected_error in result.stderr
